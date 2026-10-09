from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

OUT = r'D:\Suivi_Clients\prospects.xlsx'
NROWS = 200
FONT='Arial'; NAVY='1F4E79'; ALT='F4F8FC'
thin=Side(style='thin',color='BFBFBF'); border=Border(left=thin,right=thin,top=thin,bottom=thin)
DATEF='DD/MM/YYYY'

wb=Workbook(); ws=wb.active; ws.title='Prospects'
headers=['Date ajout','Entreprise','Secteur','URL site actuel','Faiblesses du site',
         'Coordonnées (mail/tél)','Fourchette tarif','Priorité','Statut','Lien maquette',
         'Date dernier contact','Notes']
for j,h in enumerate(headers,1):
    c=ws.cell(1,j,h)
    c.font=Font(name=FONT,bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor=NAVY)
    c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True); c.border=border
ws.row_dimensions[1].height=30
widths=[13,26,18,34,34,24,16,12,40,40,17,30]
for j,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(j)].width=w

for r in range(2,NROWS+2):
    for j in range(1,13):
        c=ws.cell(r,j); c.font=Font(name=FONT,size=10); c.border=border
        if r%2==0: c.fill=PatternFill('solid',fgColor=ALT)
    ws.cell(r,1).number_format=DATEF
    ws.cell(r,11).number_format=DATEF

dv_stat=DataValidation(type='list',formula1='"Refonte à faire,À contacter,Relance,Intéressé,Converti,Pas intéressé,Sans réponse"',allow_blank=True)
dv_prio=DataValidation(type='list',formula1='"Haute,Moyenne,Basse"',allow_blank=True)
ws.add_data_validation(dv_stat); ws.add_data_validation(dv_prio)
dv_stat.add(f'I2:I{NROWS+1}'); dv_prio.add(f'H2:H{NROWS+1}')
ws.freeze_panes='A2'

# Légende du pipeline sur une 2e feuille
lg=wb.create_sheet('Pipeline')
lg.column_dimensions['A'].width=24; lg.column_dimensions['B'].width=60
lg['A1']='PIPELINE — circulation du statut'; lg['A1'].font=Font(name=FONT,bold=True,size=14,color=NAVY)
rows=[('Statut','Posé par / signification'),
      ('Refonte à faire','Prompt PROSPECTION : prospect trouvé et enregistré.'),
      ('À contacter','Prompt REFONTE : maquette faite + déployée (lien rempli).'),
      ('Relance','Prompt CONTACT : proposition envoyée au client.'),
      ('Intéressé','Le client a répondu positivement.'),
      ('Converti','Devient client payant → bascule dans Suivi_clients.xlsx (URSSAF).'),
      ('Pas intéressé','Refus.'),
      ('Sans réponse','Aucune réponse après relances.')]
for i,(a,b) in enumerate(rows):
    r=3+i
    ca=lg.cell(r,1,a); cb=lg.cell(r,2,b)
    if i==0:
        for c in (ca,cb): c.font=Font(name=FONT,bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor=NAVY); c.border=border
    else:
        ca.font=Font(name=FONT,bold=True,size=10); cb.font=Font(name=FONT,size=10)
        ca.border=border; cb.border=border

wb.save(OUT); print('Saved',OUT)
